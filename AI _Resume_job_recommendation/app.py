from flask import Flask, render_template, request, redirect, session
import os
import re

from pypdf import PdfReader
from docx import Document


# =========================================================
# FLASK APP
# =========================================================

app = Flask(__name__)

app.secret_key = "resumeai_secret_key_2026"

UPLOAD_FOLDER = "uploads"

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)

app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER


# =========================================================
# JOB DATABASE
# =========================================================

jobs_data = [

    {
        "id": 1,
        "title": "Machine Learning Engineer",
        "company": "AI Technologies",
        "location": "Bangalore",
        "salary": "₹6 - ₹12 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "Python", "Machine Learning",
            "Scikit-learn", "Pandas", "NumPy"
        ],
        "description":
            "Develop and deploy machine learning models for real-world applications.",
        "education":
            "B.Tech / B.E / M.Tech in Computer Science or IT"
    },

    {
        "id": 2,
        "title": "Data Analyst",
        "company": "Data Insights Pvt Ltd",
        "location": "Chennai",
        "salary": "₹4 - ₹8 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "Python", "SQL", "Excel",
            "Pandas", "Data Visualization"
        ],
        "description":
            "Analyze business data and generate useful insights for decision making.",
        "education":
            "Bachelor's Degree in IT, CS, Statistics or related field"
    },

    {
        "id": 3,
        "title": "Data Scientist",
        "company": "Analytics Solutions",
        "location": "Hyderabad",
        "salary": "₹7 - ₹15 LPA",
        "experience": "1 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Python", "Machine Learning",
            "Statistics", "SQL", "Pandas"
        ],
        "description":
            "Build predictive models and analyze large datasets.",
        "education":
            "B.Tech / M.Tech / M.Sc"
    },

    {
        "id": 4,
        "title": "AI Engineer",
        "company": "Future AI Labs",
        "location": "Bangalore",
        "salary": "₹7 - ₹14 LPA",
        "experience": "0 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Python", "AI", "Machine Learning",
            "Deep Learning", "TensorFlow"
        ],
        "description":
            "Develop artificial intelligence solutions using modern AI technologies.",
        "education":
            "B.Tech / M.Tech in Computer Science or AI"
    },

    {
        "id": 5,
        "title": "Python Developer",
        "company": "TechSoft Solutions",
        "location": "Chennai",
        "salary": "₹4 - ₹9 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "Python", "Flask", "Django",
            "SQL", "REST API"
        ],
        "description":
            "Develop backend applications and APIs using Python.",
        "education":
            "B.Tech / B.E / BCA / MCA"
    },

    {
        "id": 6,
        "title": "Deep Learning Engineer",
        "company": "NeuralTech",
        "location": "Pune",
        "salary": "₹8 - ₹16 LPA",
        "experience": "1 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Python", "TensorFlow", "Keras",
            "CNN", "RNN"
        ],
        "description":
            "Build deep learning models for computer vision and NLP applications.",
        "education":
            "B.Tech / M.Tech"
    },

    {
        "id": 7,
        "title": "Data Engineer",
        "company": "Cloud Data Systems",
        "location": "Hyderabad",
        "salary": "₹6 - ₹13 LPA",
        "experience": "1 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Python", "SQL", "ETL",
            "Spark", "Databases"
        ],
        "description":
            "Build and maintain data pipelines and data infrastructure.",
        "education":
            "B.Tech / B.E / MCA"
    },

    {
        "id": 8,
        "title": "Business Analyst",
        "company": "Business Solutions Ltd",
        "location": "Mumbai",
        "salary": "₹5 - ₹10 LPA",
        "experience": "0 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Excel", "SQL", "Power BI",
            "Communication", "Analytics"
        ],
        "description":
            "Analyze business requirements and provide data-driven solutions.",
        "education":
            "Bachelor's Degree"
    },

    {
        "id": 9,
        "title": "AI/ML Intern",
        "company": "Innovate AI",
        "location": "Chennai",
        "salary": "₹15,000 - ₹30,000 / Month",
        "experience": "Fresher",
        "type": "Internship",
        "skills": [
            "Python", "Machine Learning",
            "Pandas", "NumPy"
        ],
        "description":
            "Assist AI and machine learning teams in developing models.",
        "education":
            "B.Tech / B.E / BCA / MCA"
    },

    {
        "id": 10,
        "title": "NLP Engineer",
        "company": "Language AI Systems",
        "location": "Bangalore",
        "salary": "₹7 - ₹14 LPA",
        "experience": "1 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Python", "NLP", "NLTK",
            "Transformers", "Deep Learning"
        ],
        "description":
            "Develop natural language processing and text analysis solutions.",
        "education":
            "B.Tech / M.Tech"
    },

    {
        "id": 11,
        "title": "Computer Vision Engineer",
        "company": "Vision Technologies",
        "location": "Bangalore",
        "salary": "₹7 - ₹15 LPA",
        "experience": "1 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Python", "OpenCV", "CNN",
            "TensorFlow", "Computer Vision"
        ],
        "description":
            "Develop computer vision models for image and video applications.",
        "education":
            "B.Tech / M.Tech"
    },

    {
        "id": 12,
        "title": "Software Developer",
        "company": "CodeWorks India",
        "location": "Chennai",
        "salary": "₹4 - ₹10 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "Java", "Python", "SQL",
            "HTML", "CSS"
        ],
        "description":
            "Design, develop and maintain software applications.",
        "education":
            "B.Tech / B.E / MCA"
    },

    {
        "id": 13,
        "title": "Frontend Developer",
        "company": "WebTech Solutions",
        "location": "Coimbatore",
        "salary": "₹4 - ₹9 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "HTML", "CSS",
            "JavaScript", "React"
        ],
        "description":
            "Build responsive and user-friendly web interfaces.",
        "education":
            "B.Tech / BCA / MCA"
    },

    {
        "id": 14,
        "title": "Backend Developer",
        "company": "CloudSoft",
        "location": "Hyderabad",
        "salary": "₹5 - ₹11 LPA",
        "experience": "0 - 3 Years",
        "type": "Full Time",
        "skills": [
            "Python", "Flask", "Node.js",
            "SQL", "APIs"
        ],
        "description":
            "Develop scalable backend services and APIs.",
        "education":
            "B.Tech / B.E / MCA"
    },

    {
        "id": 15,
        "title": "SQL Developer",
        "company": "Database Systems",
        "location": "Chennai",
        "salary": "₹4 - ₹8 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "SQL", "MySQL", "Database",
            "Queries", "DBMS"
        ],
        "description":
            "Develop database queries and maintain database systems.",
        "education":
            "B.Tech / B.E / BCA / MCA"
    },

    {
        "id": 16,
        "title": "Power BI Developer",
        "company": "Business Intelligence Corp",
        "location": "Bangalore",
        "salary": "₹5 - ₹10 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "Power BI", "SQL",
            "Excel", "Data Visualization"
        ],
        "description":
            "Create interactive dashboards and business intelligence reports.",
        "education":
            "Bachelor's Degree"
    },

    {
        "id": 17,
        "title": "Cloud Engineer",
        "company": "CloudWorks",
        "location": "Pune",
        "salary": "₹6 - ₹13 LPA",
        "experience": "1 - 3 Years",
        "type": "Full Time",
        "skills": [
            "AWS", "Azure", "Linux",
            "Python", "Networking"
        ],
        "description":
            "Deploy and maintain cloud-based infrastructure and services.",
        "education":
            "B.Tech / B.E"
    },

    {
        "id": 18,
        "title": "Research Analyst",
        "company": "Research Analytics",
        "location": "Chennai",
        "salary": "₹4 - ₹8 LPA",
        "experience": "0 - 2 Years",
        "type": "Full Time",
        "skills": [
            "Python", "Excel", "Statistics",
            "SQL", "Research"
        ],
        "description":
            "Conduct research and analyze datasets to identify useful patterns.",
        "education":
            "Bachelor's / Master's Degree"
    }
]


# =========================================================
# SKILL LIST
# =========================================================

KNOWN_SKILLS = [
    "Python", "Java", "C", "C++", "SQL",
    "HTML", "CSS", "JavaScript", "React",
    "Flask", "Django", "Node.js", "REST API",
    "APIs", "Machine Learning", "Deep Learning",
    "Artificial Intelligence", "AI", "Data Science",
    "Data Analysis", "Data Visualization", "Pandas",
    "NumPy", "Scikit-learn", "TensorFlow", "Keras",
    "CNN", "RNN", "NLP", "NLTK", "Transformers",
    "OpenCV", "Computer Vision", "Statistics",
    "Excel", "Power BI", "AWS", "Azure",
    "Linux", "Networking", "Spark", "ETL",
    "Databases", "Database", "MySQL", "DBMS",
    "Git", "GitHub"
]


# =========================================================
# NAME VALIDATION
# =========================================================

def is_valid_name(name):

    name = name.strip()

    if len(name) < 2:
        return False

    if len(name) > 50:
        return False

    if re.search(r"\d", name):
        return False

    if "@" in name:
        return False

    if "http" in name.lower():
        return False

    if "www." in name.lower():
        return False

    if not re.match(r"^[A-Za-z][A-Za-z .'-]*$", name):
        return False

    if len(name.split()) > 4:
        return False

    blocked = [
        "resume", "curriculum vitae", "profile",
        "summary", "objective", "education",
        "skills", "experience", "projects",
        "project", "training", "internship",
        "developer", "engineer", "analyst",
        "technology", "technologies",
        "machine learning", "data science",
        "artificial intelligence", "professional",
        "responsibilities", "certification",
        "currently", "working", "knowledge",
        "proficient", "using", "developed",
        "application", "applications"
    ]

    lower_name = name.lower()

    for word in blocked:
        if word in lower_name:
            return False

    if name.endswith("."):
        return False

    return True


# =========================================================
# EXTRACT CANDIDATE NAME
# =========================================================

def extract_candidate_name(text):

    if not text:
        return "Candidate"

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    # Name: XXXXX
    for line in lines:

        match = re.search(
            r"^(?:name|candidate\s+name|full\s+name)"
            r"\s*[:\-]\s*(.+)$",
            line,
            re.IGNORECASE
        )

        if match:

            name = match.group(1).strip()

            if is_valid_name(name):
                return name

    # Look at first few lines
    for line in lines[:12]:

        clean = line.strip()
        lower = clean.lower()

        if "@" in clean:
            continue

        if "http" in lower:
            continue

        if "linkedin" in lower:
            continue

        if "github" in lower:
            continue

        if re.search(r"\d", clean):
            continue

        if len(clean.split()) > 4:
            continue

        if is_valid_name(clean):
            return clean

    return "Candidate"


# =========================================================
# EXTRACT EDUCATION
# =========================================================

def extract_education(text):

    lines = text.splitlines()

    education_patterns = [

        r"\bbachelor\s+of\s+computer\s+applications\b",
        r"\bbca\b",

        r"\bbachelor\s+of\s+technology\b",
        r"\bb\.?\s*tech\b",
        r"\bb\.?\s*e\b",

        r"\bbachelor\s+of\s+science\b",
        r"\bb\.?\s*sc\b",

        r"\bmaster\s+of\s+computer\s+applications\b",
        r"\bmca\b",

        r"\bmaster\s+of\s+technology\b",
        r"\bm\.?\s*tech\b",

        r"\bmaster\s+of\s+science\b",
        r"\bm\.?\s*sc\b",

        r"\bmba\b",
        r"\bph\.?\s*d\b",
        r"\bphd\b"
    ]

    # First priority:
    # Bachelor of Computer Applications
    for line in lines:

        clean = line.strip()

        if re.search(
            r"bachelor\s+of\s+computer\s+applications",
            clean,
            re.IGNORECASE
        ):

            return "Bachelor of Computer Applications"

    # BCA
    for line in lines:

        clean = line.strip()

        if re.search(
            r"\bbca\b",
            clean,
            re.IGNORECASE
        ):

            return "Bachelor of Computer Applications"

    # Other education
    for line in lines:

        clean = line.strip()

        if len(clean) > 150:
            continue

        for pattern in education_patterns:

            if re.search(
                pattern,
                clean,
                re.IGNORECASE
            ):

                # Do not return summary paragraphs
                if len(clean.split()) <= 15:
                    return clean

    return "Not detected"


# =========================================================
# EXTRACT EXPERIENCE
# =========================================================

def extract_experience(text):

    lower = text.lower()

    patterns = [

        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+of\s+experience",

        r"(\d+(?:\.\d+)?)\s*\+?\s*years?\s+experience",

        r"experience\s*[:\-]\s*(.{1,80})",

    ]

    for pattern in patterns:

        match = re.search(
            pattern,
            lower,
            re.IGNORECASE
        )

        if match:

            value = match.group(1).strip()

            if value:

                if value.replace(".", "", 1).isdigit():
                    return value + " Years"

                return value

    if "fresher" in lower:
        return "Fresher"

    if "intern" in lower or "internship" in lower:
        return "Internship / Entry Level"

    return "Not detected"


# =========================================================
# EXTRACT SKILLS
# =========================================================

def extract_skills(text):

    lower_text = text.lower()

    detected = []

    for skill in KNOWN_SKILLS:

        pattern = re.escape(skill.lower())

        if re.search(
            r"(?<![a-z])" +
            pattern +
            r"(?![a-z])",
            lower_text
        ):

            if skill not in detected:
                detected.append(skill)

    return detected


# =========================================================
# EXTRACT CAREER FIELD
# =========================================================

def extract_career_field(text, skills):

    lower = text.lower()

    if (
        "machine learning" in lower
        or "deep learning" in lower
        or "tensorflow" in lower
        or "scikit-learn" in lower
    ):
        return "Machine Learning / AI"

    if (
        "data scientist" in lower
        or "data science" in lower
        or "statistics" in lower
    ):
        return "Data Science"

    if (
        "data analyst" in lower
        or "data analysis" in lower
        or "power bi" in lower
    ):
        return "Data Analytics"

    if (
        "frontend" in lower
        or "react" in lower
        or "html" in lower
        or "css" in lower
    ):
        return "Web Development"

    if (
        "backend" in lower
        or "flask" in lower
        or "django" in lower
        or "api" in lower
    ):
        return "Backend Development"

    if (
        "cloud" in lower
        or "aws" in lower
        or "azure" in lower
    ):
        return "Cloud Computing"

    if (
        "java" in lower
        or "software developer" in lower
    ):
        return "Software Development"

    if skills:
        return "Information Technology"

    return "Not detected"


# =========================================================
# RESUME SCORE
# =========================================================

def calculate_resume_score(
    text,
    name,
    education,
    experience,
    skills
):

    score = 0

    if name != "Candidate":
        score += 20

    if education != "Not detected":
        score += 20

    if experience != "Not detected":
        score += 15

    if len(skills) >= 5:
        score += 25

    elif len(skills) >= 3:
        score += 20

    elif len(skills) >= 1:
        score += 10

    if len(text) > 500:
        score += 10

    if "project" in text.lower() or "projects" in text.lower():
        score += 5

    if (
        "certification" in text.lower()
        or "certifications" in text.lower()
        or "certificate" in text.lower()
    ):
        score += 5

    return min(score, 100)


# =========================================================
# READ PDF
# =========================================================

def extract_pdf_text(file_path):

    text = ""

    reader = PdfReader(file_path)

    for page in reader.pages:

        page_text = page.extract_text()

        if page_text:
            text += page_text + "\n"

    return text


# =========================================================
# READ DOCX
# =========================================================

def extract_docx_text(file_path):

    text = ""

    document = Document(file_path)

    for paragraph in document.paragraphs:

        if paragraph.text.strip():
            text += paragraph.text + "\n"

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():
                    text += cell.text + "\n"

    return text


# =========================================================
# ANALYZE RESUME
# =========================================================

def analyze_resume(text):

    name = extract_candidate_name(text)

    education = extract_education(text)

    experience = extract_experience(text)

    skills = extract_skills(text)

    career_field = extract_career_field(
        text,
        skills
    )

    score = calculate_resume_score(
        text,
        name,
        education,
        experience,
        skills
    )

    return {
        "name": name,
        "education": education,
        "experience": experience,
        "skills": skills,
        "career_field": career_field,
        "score": score
    }


# =========================================================
# JOB MATCHING
# =========================================================

def get_recommended_jobs(user_skills):

    user_skills_lower = [
        skill.lower()
        for skill in user_skills
    ]

    recommendations = []

    for job in jobs_data:

        job_skills = [
            skill.lower()
            for skill in job["skills"]
        ]

        matched = []

        for skill in job_skills:

            if skill in user_skills_lower:
                matched.append(skill)

        if job_skills:

            match_percentage = round(
                (len(matched) / len(job_skills)) * 100
            )

        else:

            match_percentage = 0

        job_copy = job.copy()

        job_copy["matched_skills"] = matched

        job_copy["match_percentage"] = match_percentage

        recommendations.append(job_copy)

    recommendations.sort(
        key=lambda x: x["match_percentage"],
        reverse=True
    )

    return recommendations


# =========================================================
# SKILL GAP
# =========================================================

def get_skill_gap(user_skills):

    user_lower = [
        skill.lower()
        for skill in user_skills
    ]

    all_job_skills = []

    for job in jobs_data:

        for skill in job["skills"]:

            if skill.lower() not in all_job_skills:
                all_job_skills.append(skill.lower())

    missing = []

    for skill in all_job_skills:

        if skill not in user_lower:

            original = next(
                (
                    s
                    for job in jobs_data
                    for s in job["skills"]
                    if s.lower() == skill
                ),
                skill
            )

            missing.append(original)

    return missing[:15]


# =========================================================
# HOME
# =========================================================

@app.route("/")
def home():

    return render_template("index.html")


# =========================================================
# ABOUT
# =========================================================

@app.route("/about")
def about():

    return render_template("about.html")


# =========================================================
# UPLOAD
# =========================================================

@app.route(
    "/upload",
    methods=["GET", "POST"]
)
def upload():

    if request.method == "POST":

        file = request.files.get("resume")

        if not file or not file.filename:

            return "Please select a resume."

        extension = os.path.splitext(
            file.filename
        )[1].lower()

        if extension not in [".pdf", ".docx"]:

            return (
                "Only PDF and DOCX files "
                "are supported."
            )

        # Delete previous uploaded files
        for old_file in os.listdir(
            app.config["UPLOAD_FOLDER"]
        ):

            old_path = os.path.join(
                app.config["UPLOAD_FOLDER"],
                old_file
            )

            try:

                if os.path.isfile(old_path):
                    os.remove(old_path)

            except Exception:
                pass

        filename = os.path.basename(
            file.filename
        )

        file_path = os.path.join(
            app.config["UPLOAD_FOLDER"],
            filename
        )

        file.save(file_path)

        try:

            if extension == ".pdf":

                resume_text = extract_pdf_text(
                    file_path
                )

            else:

                resume_text = extract_docx_text(
                    file_path
                )

        except Exception as e:

            return (
                "Unable to read resume: "
                + str(e)
            )

        if not resume_text.strip():

            return (
                "Could not extract text from "
                "this resume. Please upload "
                "a text-based PDF or DOCX."
            )

        # Analyze resume
        data = analyze_resume(
            resume_text
        )

        # =================================================
        # RESUME HISTORY
        # =================================================

        # Get old history BEFORE changing current session
        old_history = session.get(
            "resume_history",
            []
        )

        # Make a new history list
        resume_history = list(old_history)

        # Add current resume
        resume_history.append({
            "candidate_name": data["name"],
            "resume_filename": filename,
            "resume_score": data["score"],
            "skills": data["skills"],
            "education": data["education"],
            "experience": data["experience"],
            "career_field": data["career_field"]
        })

        # Keep only latest 20 resumes
        resume_history = resume_history[-20:]

        # =================================================
        # CURRENT RESUME DATA
        # =================================================

        session["resume_history"] = resume_history

        session["candidate_name"] = data["name"]

        session["education"] = data["education"]

        session["experience"] = data["experience"]

        session["skills"] = data["skills"]

        session["career_field"] = data["career_field"]

        session["resume_score"] = data["score"]

        session["resume_filename"] = filename

        session.modified = True

        # Recommendations
        recommendations = get_recommended_jobs(
            data["skills"]
        )

        session["recommended_job_ids"] = [
            job["id"]
            for job in recommendations
        ]

        # Skill gap
        session["skill_gap"] = get_skill_gap(
            data["skills"]
        )

        print("\n==============================")
        print("NEW RESUME UPLOADED")
        print("==============================")
        print("Name:", data["name"])
        print("Education:", data["education"])
        print("Experience:", data["experience"])
        print("Career:", data["career_field"])
        print("Skills:", data["skills"])
        print("Score:", data["score"])
        print("File:", filename)
        print("History Count:", len(resume_history))
        print("==============================\n")

        return redirect("/analysis")

    return render_template(
        "upload.html"
    )


# =========================================================
# ANALYSIS
# =========================================================

@app.route("/analysis")
def analysis():

    if "resume_filename" not in session:
        return redirect("/upload")

    return render_template(
        "analysis.html",

        candidate_name=session.get(
            "candidate_name",
            ""
        ),

        education=session.get(
            "education",
            ""
        ),

        experience=session.get(
            "experience",
            ""
        ),

        skills=session.get(
            "skills",
            []
        ),

        career_field=session.get(
            "career_field",
            ""
        ),

        resume_score=session.get(
            "resume_score",
            0
        ),

        resume_filename=session.get(
            "resume_filename",
            ""
        )
    )


# =========================================================
# JOBS
# =========================================================

@app.route("/jobs")
def jobs():

    if "resume_filename" not in session:
        return redirect("/upload")

    user_skills = session.get(
        "skills",
        []
    )

    recommendations = get_recommended_jobs(
        user_skills
    )

    return render_template(
        "jobs.html",

        jobs=recommendations,

        candidate_name=session.get(
            "candidate_name",
            ""
        ),

        skills=user_skills
    )


# =========================================================
# JOB DETAILS
# =========================================================

@app.route(
    "/job-details/<int:job_id>"
)
def job_details(job_id):

    selected_job = None

    for job in jobs_data:

        if job["id"] == job_id:

            selected_job = job
            break

    if selected_job is None:
        return "Job not found", 404

    user_skills = session.get(
        "skills",
        []
    )

    user_lower = [
        skill.lower()
        for skill in user_skills
    ]

    matched = []

    for skill in selected_job["skills"]:

        if skill.lower() in user_lower:
            matched.append(skill)

    if selected_job["skills"]:

        match_percentage = round(
            len(matched)
            / len(selected_job["skills"])
            * 100
        )

    else:

        match_percentage = 0

    return render_template(
        "job_details.html",

        job=selected_job,

        matched_skills=matched,

        match_percentage=match_percentage,

        candidate_name=session.get(
            "candidate_name",
            ""
        )
    )


# =========================================================
# SKILL GAP
# =========================================================

@app.route("/skill-gap")
def skill_gap():

    if "resume_filename" not in session:
        return redirect("/upload")

    user_skills = session.get(
        "skills",
        []
    )

    missing_skills = get_skill_gap(
        user_skills
    )

    return render_template(
        "skill_gap.html",

        candidate_name=session.get(
            "candidate_name",
            ""
        ),

        user_skills=user_skills,

        missing_skills=missing_skills
    )


# =========================================================
# DASHBOARD
# =========================================================

@app.route("/dashboard")
def dashboard():

    if "resume_filename" not in session:
        return redirect("/upload")

    user_skills = session.get(
        "skills",
        []
    )

    recommendations = get_recommended_jobs(
        user_skills
    )

    return render_template(
        "dashboard.html",

        candidate_name=session.get(
            "candidate_name",
            ""
        ),

        education=session.get(
            "education",
            ""
        ),

        experience=session.get(
            "experience",
            ""
        ),

        career_field=session.get(
            "career_field",
            ""
        ),

        skills=user_skills,

        resume_score=session.get(
            "resume_score",
            0
        ),

        jobs=recommendations[:5]
    )


# =========================================================
# CAREER INSIGHTS
# =========================================================

@app.route("/career-insights")
def career_insights():

    if "resume_filename" not in session:
        return redirect("/upload")

    user_skills = session.get(
        "skills",
        []
    )

    recommendations = get_recommended_jobs(
        user_skills
    )

    return render_template(
        "career_insights.html",

        candidate_name=session.get(
            "candidate_name",
            ""
        ),

        career_field=session.get(
            "career_field",
            ""
        ),

        skills=user_skills,

        resume_score=session.get(
            "resume_score",
            0
        ),

        jobs=recommendations
    )


# =========================================================
# RESUME HISTORY
# =========================================================

@app.route("/resume-history")
def resume_history():

    if "resume_history" not in session:
        return redirect("/upload")

    history = session.get(
        "resume_history",
        []
    )

    return render_template(
        "resume_history.html",

        history=history,

        candidate_name=session.get(
            "candidate_name",
            ""
        ),

        resume_filename=session.get(
            "resume_filename",
            ""
        ),

        resume_score=session.get(
            "resume_score",
            0
        ),

        skills=session.get(
            "skills",
            []
        )
    )


# =========================================================
# SAVED JOBS
# =========================================================

@app.route("/saved-jobs")
def saved_jobs():

    if "resume_filename" not in session:
        return redirect("/upload")

    saved_job_ids = session.get(
        "recommended_job_ids",
        []
    )

    saved_jobs_list = []

    for job_id in saved_job_ids:

        for job in jobs_data:

            if job["id"] == job_id:

                job_copy = job.copy()

                user_skills = session.get(
                    "skills",
                    []
                )

                user_skills_lower = [
                    skill.lower()
                    for skill in user_skills
                ]

                matched = []

                for skill in job["skills"]:

                    if skill.lower() in user_skills_lower:
                        matched.append(skill)

                if job["skills"]:

                    match_percentage = round(
                        len(matched)
                        / len(job["skills"])
                        * 100
                    )

                else:

                    match_percentage = 0

                job_copy["matched_skills"] = matched

                job_copy["match_percentage"] = match_percentage

                saved_jobs_list.append(
                    job_copy
                )

                break

    return render_template(
        "saved_jobs.html",

        jobs=saved_jobs_list,

        candidate_name=session.get(
            "candidate_name",
            ""
        )
    )


# =========================================================
# PROFILE
# =========================================================

@app.route("/profile")
def profile():

    if "resume_filename" not in session:
        return redirect("/upload")

    return render_template(
        "profile.html",

        candidate_name=session.get(
            "candidate_name",
            ""
        ),

        education=session.get(
            "education",
            ""
        ),

        experience=session.get(
            "experience",
            ""
        ),

        career_field=session.get(
            "career_field",
            ""
        ),

        skills=session.get(
            "skills",
            []
        ),

        resume_score=session.get(
            "resume_score",
            0
        ),

        resume_filename=session.get(
            "resume_filename",
            ""
        )
    )


# =========================================================
# CLEAR CURRENT RESUME
# =========================================================

@app.route("/clear-resume")
def clear_resume():

    # Preserve history
    history = session.get(
        "resume_history",
        []
    )

    session.clear()

    # Restore history
    session["resume_history"] = history

    session.modified = True

    return redirect("/upload")



# =========================================================
# RUN SERVER
# =========================================================

if __name__ == "__main__":

    app.run(
        debug=True
    )